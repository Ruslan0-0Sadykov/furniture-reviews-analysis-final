"""Fail closed when the Git publication set contains secrets or review data."""

from __future__ import annotations

import csv
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PRIVATE_REVIEWS = ROOT / "data" / "processed" / "reviews_enriched.csv"
BLOCKED_PATH_PARTS = (
    "data/processed/",
    "classification_audit",
    "h5_price_records",
    "datalens/datalens_dataset.csv",
    "outputs/metrics/hypotheses.json",
)
BLOCKED_PUBLIC_COLUMNS = {
    "author",
    "author_status",
    "profile",
    "profile_url",
    "review_id",
    "review_text",
    "text",
    "text_original",
    "address",
    "coordinates",
    "company",
}
SECRET_PATTERNS = {
    "private key": re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
    "GitHub token": re.compile(r"\b(?:gh[opsu]_[A-Za-z0-9]{30,}|github_pat_[A-Za-z0-9_]{30,})\b"),
    "AWS access key": re.compile(r"\bAKIA[0-9A-Z]{16}\b"),
    "JWT": re.compile(r"\beyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\b"),
    "Yandex API key": re.compile(r"\bAQVN[A-Za-z0-9_-]{20,}\b"),
}


def tracked_paths() -> list[Path]:
    raw = subprocess.check_output(
        ["git", "ls-files", "-z"], cwd=ROOT
    ).decode("utf-8", errors="strict")
    return [ROOT / item for item in raw.split("\0") if item]


def visible_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".csv", ".json", ".py", ".ipynb", ".yml", ".yaml"}:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix == ".docx":
        doc = Document(path)
        parts = [paragraph.text for paragraph in doc.paragraphs]
        parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        for section in doc.sections:
            parts.extend(paragraph.text for paragraph in section.header.paragraphs)
            parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
        return "\n".join(parts)
    if suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        return "\n".join(
            str(cell.value)
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    return ""


def normalized(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def env_secret_values() -> list[str]:
    env_path = ROOT / ".env"
    if not env_path.exists():
        return []
    values: list[str] = []
    for line in env_path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        if not line or line.lstrip().startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip("\"'")
        if re.search(r"KEY|TOKEN|SECRET|PASSWORD", key, re.I) and len(value) >= 8:
            if "your_" not in value.casefold() and "example" not in value.casefold():
                values.append(value)
    return values


def main() -> int:
    paths = tracked_paths()
    relative = [path.relative_to(ROOT).as_posix() for path in paths]
    errors: list[str] = []

    for name in relative:
        if any(part in name for part in BLOCKED_PATH_PARTS):
            errors.append(f"blocked path tracked: {name}")

    public_dataset = ROOT / "datalens" / "datalens_dataset_public.csv"
    with public_dataset.open(encoding="utf-8-sig", newline="") as stream:
        header = set(next(csv.reader(stream)))
    blocked_columns = sorted(header & BLOCKED_PUBLIC_COLUMNS)
    if blocked_columns:
        errors.append("blocked public columns: " + ", ".join(blocked_columns))

    texts: dict[str, str] = {}
    for path, name in zip(paths, relative):
        text = visible_text(path)
        texts[name] = text
        for label, pattern in SECRET_PATTERNS.items():
            if pattern.search(text):
                errors.append(f"{label} pattern: {name}")

    actual_secrets = env_secret_values()
    for secret in actual_secrets:
        encoded = secret.encode("utf-8")
        for path, name in zip(paths, relative):
            if encoded in path.read_bytes():
                errors.append(f"local .env secret value: {name}")

    private = pd.read_csv(
        PRIVATE_REVIEWS, usecols=["author", "text_original"], low_memory=False
    )
    authors = {
        normalized(value)
        for value in private["author"].dropna()
        if len(normalized(value)) >= 6
    }
    review_prefixes = {
        normalized(value)[:80]
        for value in private["text_original"].dropna()
        if len(normalized(value)) >= 80
    }
    for name, text in texts.items():
        # The repository locator necessarily contains the owner's public
        # GitHub handle; it is not a value exported from the review corpus.
        if name == "GITHUB_URL.txt":
            continue
        haystack = normalized(text)
        if any(author in haystack for author in authors):
            errors.append(f"author value match: {name}")
        if any(prefix in haystack for prefix in review_prefixes):
            errors.append(f"review text prefix match: {name}")

    if errors:
        print("PUBLIC AUDIT: FAILED")
        for error in sorted(set(errors)):
            print(f"- {error}")
        return 1

    print(f"PUBLIC AUDIT: PASSED ({len(paths)} tracked files)")
    print(f"Actual local secret values checked without disclosure: {len(actual_secrets)}")
    print("Review author matches: 0")
    print("Review text prefix matches: 0")
    print("Blocked public dataset columns: 0")
    return 0


if __name__ == "__main__":
    sys.exit(main())

from __future__ import annotations

import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)


def _add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"
    for i,h in enumerate(headers):
        cell=table.rows[0].cells[i]; cell.text=str(h); _set_cell_shading(cell,"1F4E78"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(8)
    trPr=table.rows[0]._tr.get_or_add_trPr(); tblHeader=OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"),"true"); trPr.append(tblHeader)
    for r_idx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text="" if value is None else str(value); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx%2: _set_cell_shading(cells[i],"F3F6FA")
            for p in cells[i].paragraphs:
                for run in p.runs: run.font.size=Pt(8)
    for row_index,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index < len(table.rows)-1:
            for cell in row.cells:
                for paragraph in cell.paragraphs: paragraph.paragraph_format.keep_with_next=True
    return table


def _pct(value): return "н.д." if value is None else f"{100*value:.1f}%"
def _num(value,digits=3): return "н.д." if value is None else f"{value:.{digits}f}"


def create_report(df, audit, results, figures, output_path: Path):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.0)
    styles=doc.styles
    styles["Normal"].font.name="Arial"; styles["Normal"].font.size=Pt(10.5); styles["Normal"].paragraph_format.space_after=Pt(6); styles["Normal"].paragraph_format.line_spacing=1.15
    for name,size in [("Title",20),("Heading 1",15),("Heading 2",12)]:
        st=styles[name]; st.font.name="Arial"; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0); st.font.bold=True
        pPr=st.element.get_or_add_pPr()
        for child in list(pPr):
            if child.tag == qn("w:pBdr"): pPr.remove(child)
    p=doc.add_paragraph(style="Title"); p.add_run("Аналитический отчет по отзывам покупателей мебели")
    pPr=p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn("w:pBdr"): pPr.remove(child)
    p=doc.add_paragraph(); p.add_run("Производственная практика").bold=True
    doc.add_paragraph(f"Исследование охватывает {len(df):,} очищенных записей из трех открытых источников. Рейтинг используется как независимый критерий тональности, а тематические признаки извлекаются из текста прозрачными правилами и, при доступности, ограниченной автоматической разметкой YandexGPT.")
    doc.add_heading("1 Введение",level=1); doc.add_paragraph("Отзывы покупателей содержат прямые сигналы о качестве товара, доставке, работе сотрудников и претензионном опыте. Анализ нужен для выделения повторяющихся потребностей и проверки пяти заранее сформулированных бизнес-гипотез.")
    doc.add_heading("2 Цель проекта",level=1); doc.add_paragraph("Цель проекта — определить ключевые потребности и болевые точки покупателей мебели, оценить статистическую поддержку пяти гипотез и перевести результаты в измеримые рекомендации для компании.")
    doc.add_heading("3 Описание исходных данных",level=1)
    _add_table(doc,["Показатель","Значение"],[["Строк загружено",audit["preprocessing"]["rows_loaded"]],["Строк после очистки",audit["preprocessing"]["rows_processed"]],["Отзывы с рейтингом",audit["preprocessing"]["rating_nonmissing"]],["Отзывы с текстом",audit["preprocessing"]["text_nonempty"]],["Удалено технических строк",audit["preprocessing"]["blank_or_nonreview_rows_removed"]],["Удалено точных дублей",audit["preprocessing"]["exact_duplicates_removed"]]])
    doc.add_picture(figures["rating"],width=Cm(15.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("4 Подготовка и очистка данных",level=1); doc.add_paragraph("Названия полей унифицированы, рейтинги преобразованы в числа, HTML и лишние пробелы удалены. Исходный текст сохранен отдельно, поскольку знаки препинания, регистр и emoji участвуют в оценке эмоциональности. Даты приведены к календарному формату там, где исходная запись позволяла это сделать. Каждому отзыву присвоен стабильный review_id.")
    doc.add_heading("5 Первичный анализ",level=1); doc.add_picture(figures["sentiment"],width=Cm(15.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_picture(figures["complaints"],width=Cm(15.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("6 Использование YandexGPT",level=1); doc.add_paragraph("Проект реализует официальный API Yandex Cloud, строгую JSON-схему, валидацию, таймаут, повторные попытки с экспоненциальной задержкой, локальный кеш и журнал ошибок. Перед отправкой удаляются телефоны и адреса электронной почты. Основной массив размечается правилами; LLM применяется только к ограниченной выборке неоднозначных случаев. Такая разметка является автоматической, а не ручной.")
    if audit.get("llm",{}).get("successful",0):
        llm=audit["llm"]
        doc.add_paragraph(f"Минимальный тестовый запрос завершился с HTTP 200. В ограниченной стратифицированной выборке автоматически размечено {llm['successful']} отзывов из {llm.get('strata','нескольких')} страт «источник × тональность». При финальной повторной сборке {llm.get('cache_hits',0)} результатов прочитано из локального кеша без повторных API-запросов. Для остальных отзывов сохранена воспроизводимая разметка правилами.")
    elif audit.get("llm",{}).get("blocked_reason"):
        doc.add_paragraph("После проверки роли и идентификатора каталога тестовый вызов модели снова отклонен API с кодом HTTP 403 PermissionDenied для ресурсов каталога, облака или организации. Поэтому результаты исследования не используют LLM-ответы и полностью воспроизводятся локальными правилами.")
    doc.add_heading("7 Методика проверки гипотез",level=1); doc.add_paragraph("Негативным считается рейтинг 1–2, нейтральным — 3, положительным — 4–5. Для долей рассчитаны интервалы Уилсона. Непараметрические сравнения используют Mann–Whitney U и Cliff's delta. Для корреляции цены и подробности применен коэффициент Спирмена с bootstrap-интервалом. При нескольких основных тестах p-value скорректированы методом Benjamini–Hochberg.")
    figure_keys={"H1":"h1","H2":"h2","H3":"h3","H4":"h4_length","H5":"h5"}
    for idx,r in enumerate(results,8):
        if r["hypothesis"]=="H5": doc.add_page_break()
        doc.add_heading(f"{idx} Гипотеза {r['hypothesis'][1:]}",level=1)
        formulations={"H1":"Клиенты чаще оставляют негативные отзывы из-за уровня сервиса, а не из-за качества мебели.","H2":"Высокая скорость доставки связана с более высокой вероятностью положительной оценки.","H3":"Наиболее частая причина возвратов и претензий — несоответствие товара описанию.","H4":"Покупатели мягкой мебели оставляют более эмоциональные и длинные отзывы.","H5":"Чем выше цена товара, тем более подробный отзыв оставляет клиент."}
        doc.add_paragraph(formulations[r["hypothesis"]]); doc.add_paragraph(f"Выборка: n = {r['n']}. Метод: {r['test']}. Raw p-value: {_num(r.get('p_value'))}; adjusted p-value: {_num(r.get('adjusted_p_value'))}; размер эффекта: {_num(r.get('effect_size'))}.")
        if r["hypothesis"]=="H1":
            doc.add_page_break()
            m=r["metrics"]; _add_table(doc,["Причина","Количество","Доля","95% ДИ"],[["Сервис",m["service_count"],_pct(m["service_share"]),f"{_pct(m['service_ci'][0])}–{_pct(m['service_ci'][1])}"],["Качество мебели",m["quality_count"],_pct(m["quality_share"]),f"{_pct(m['quality_ci'][0])}–{_pct(m['quality_ci'][1])}"],["Обе причины",m["both_count"],"—","—"]])
        elif r["hypothesis"]=="H2":
            rows=[]
            for k,v in r["metrics"]["rates"].items(): rows.append([{"fast":"Быстро","normal":"В срок","slow":"Медленно"}.get(k,k),v["n"],v["positive"],_pct(v["share"]),f"{_pct(v['ci'][0])}–{_pct(v['ci'][1])}"])
            _add_table(doc,["Скорость","n","Положительных","Доля","95% ДИ"],rows)
        elif r["hypothesis"]=="H3":
            _add_table(doc,["Причина","Количество","Доля","95% ДИ"],[[x["reason"],x["count"],_pct(x["share"]),f"{_pct(x['ci'][0])}–{_pct(x['ci'][1])}"] for x in r["metrics"]["reason_counts"][:8]])
        elif r["hypothesis"]=="H4":
            rows=[]
            for metric,label in [("word_count","Количество слов"),("emotion_intensity","Эмоциональность")]:
                x=r["metrics"][metric]; rows.append([label,"Мягкая",x["soft"].get("n"),_num(x["soft"].get("mean")),_num(x["soft"].get("median")),_num(x["p_value"]),_num(x["cliffs_delta"])]); rows.append([label,"Корпусная",x["cabinet"].get("n"),_num(x["cabinet"].get("mean")),_num(x["cabinet"].get("median")),"—","—"])
            _add_table(doc,["Метрика","Группа","n","Среднее","Медиана","p-value","Cliff delta"],rows)
        else:
            m=r["metrics"]; _add_table(doc,["Показатель","Значение"],[["Отзывы с явной ценой",r["n"]],["Покрытие",_pct(m["price_coverage"])],["Spearman rho",_num(m["spearman_rho"])],["95% ДИ",f"{_num(m['rho_ci'][0])}–{_num(m['rho_ci'][1])}"],["Формула detail_score",m["detail_formula"]]])
        doc.add_picture(figures[figure_keys[r["hypothesis"]]],width=Cm(15.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        if r["hypothesis"]=="H4": doc.add_picture(figures["h4_emotion"],width=Cm(15.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Ограничения: {r['limitations']}")
        p=doc.add_paragraph(); p.add_run("Вердикт: ").bold=True; p.add_run(r["verdict"] + ".")
    doc.add_heading("13 Аналитический монитор",level=1); doc.add_paragraph("Для Yandex DataLens подготовлен плоский датасет, словарь полей и спецификация дашборда. Автоматическое создание монитора зависит от доступности авторизованной сессии Yandex DataLens; локальные материалы позволяют собрать его без повторной обработки данных.")
    doc.add_heading("14 Основные потребности клиентов",level=1); doc.add_paragraph("Потребности определяются по частоте и контексту упоминаний: предсказуемая доставка, корректная коммуникация, соответствие товара ожиданиям, отсутствие дефектов и понятный процесс претензии. Количественная приоритизация приведена в разделах гипотез и графике жалоб.")
    doc.add_heading("15 Основные болевые точки",level=1); doc.add_paragraph("Болевые точки сосредоточены в сервисных контактах, качестве и доставке. Их относительная значимость должна оцениваться по долям в негативных отзывах, а не по общей частоте слов, поскольку положительные отзывы также упоминают менеджеров и доставку.")
    doc.add_heading("16 Рекомендации компании",level=1)
    h1=results[0]["metrics"]; h2=results[1]["metrics"]; recs=[
        ["Негативные отзывы о сервисе",f"{h1['service_count']} из {results[0]['n']} негативных отзывов", "Разбирать обращения по этапам: консультация, доставка, сборка, гарантия; назначить владельца причины и срок ответа 24 часа.","Снижение повторных жалоб","Доля service_issue среди рейтингов 1–2; медиана времени ответа"],
        ["Непредсказуемая доставка",f"Явно классифицировано {results[1]['n']} отзывов о скорости", "Фиксировать обещанную и фактическую дату, автоматически предупреждать о переносе и измерять соблюдение обещанного интервала.","Меньше негативных оценок, связанных с задержкой","On-time delivery; P(positive) по группам скорости"],
        ["Возвраты и претензии",f"Проанализировано {results[2]['n']} отзывов с возвратом или претензией", "Добавить обязательный код причины, фото дефекта и связь с карточкой товара; ежемесячно разбирать TOP-3 причин.","Быстрее устранять повторяющиеся дефекты карточек и комплектации","Доля каждой причины; срок закрытия претензии; повторяемость по SKU"],
    ]; _add_table(doc,["Проблема","Данные","Действие","Ожидаемый эффект","KPI"],recs)
    doc.add_heading("17 Ограничения исследования",level=1); doc.add_paragraph("Данные наблюдательные и не позволяют утверждать причинность. Категория, цена, скорость доставки и причины претензий отсутствуют в исходных колонках и извлекаются из текста. Не все отзывы содержат текст; относительные даты Google восстановлены приближенно относительно даты выгрузки. Автоматическая классификация требует дополнительной человеческой валидации перед операционными решениями.")
    doc.add_heading("18 Заключение",level=1); doc.add_paragraph("Исследование дает воспроизводимую оценку пяти гипотез без подгонки результатов. Наиболее надежны выводы, основанные на рейтинге и крупных подвыборках. Результаты, зависящие от редких упоминаний цены или возврата, интерпретируются с ограничениями.")
    doc.add_page_break()
    doc.add_heading("Сводная таблица гипотез",level=1); _add_table(doc,["Гипотеза","Вердикт","Ключевые цифры","Статистическое основание","Комментарий"],[[r["hypothesis"],r["verdict"],f"n={r['n']}; effect={_num(r.get('effect_size'))}",f"p={_num(r.get('p_value'))}; adj={_num(r.get('adjusted_p_value'))}",r["limitations"]] for r in results])
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("Анализ отзывов покупателей мебели")
    output_path.parent.mkdir(parents=True,exist_ok=True); doc.save(output_path)
